"""Đọc file Google Sheets "ĐỔ DẦU - BẢO DƯỠNG - ĐĂNG KIỂM - THAY VỎ - GIẤY ĐI ĐƯỜNG".

Nguyên tắc (GĐ2):
- Google Sheet là NGUỒN GỐC. App chỉ đọc, cache ngắn (CACHE_SECONDS) + nút "Làm mới".
- Đọc theo TÊN TAB và TIÊU ĐỀ CỘT, không theo gid hay vị trí cột cố định.
- "Còn lại (ngày/km)" tự tính lại từ ngày hạn / km, không tin ô công thức (ô có thể bị gõ đè).
- Ngưỡng màu giấy tờ lấy từ chính sheet (ô "Vàng ≤", "Đỏ ≤").
"""
import io
import os
import re
import threading
import time
import unicodedata
import urllib.request
from collections import defaultdict
from datetime import date, datetime, timedelta

import openpyxl

from fleet import FLEET_BY_CLEAN, clean_plate

SHEET_ID = os.environ.get("FLEET_SHEET_ID", "1SzlPdtSjhBeqvFlZq5WiznFOVtF6SEvQ-5xKFF6I8Es")
CACHE_SECONDS = 180
FUEL_NORM = 40.0           # định mức L/100km
FUEL_DEVIATION = 0.10      # lệch > 10% thì cảnh báo
OIL_SOON_KM, OIL_NEAR_KM = 1000, 3000   # nhớt máy: còn ≤1.000 km = sắp hạn, ≤3.000 km = gần hạn

# ---------------------------------------------------------------- tiện ích


def norm(text) -> str:
    """Bỏ dấu, viết hoa, gộp khoảng trắng: 'Hạn GDD\\nĐẦU KÉO' -> 'HAN GDD DAU KEO'."""
    s = unicodedata.normalize("NFD", str(text or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").replace("Đ", "D").replace("đ", "d")
    return re.sub(r"\s+", " ", s).strip().upper()


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:  # số seri ngày kiểu Excel
        return (datetime(1899, 12, 30) + timedelta(days=float(v))).date()
    m = re.match(r"^\s*(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})", str(v))
    if m:
        try:
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


def to_num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace("đ", "")
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", s):  # 1.234.567,5
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def txt(v) -> str:
    return "" if v is None else str(v).strip()


def fmt_plate(p: str) -> str:
    c = clean_plate(p)
    return FLEET_BY_CLEAN.get(c, {}).get("plate_number") or p


def iso(d):
    return d.isoformat() if d else None


# ---------------------------------------------------------------- đọc bảng


class Table:
    """1 bảng trong 1 tab: tìm dòng tiêu đề theo tên cột, đọc các dòng dữ liệu bên dưới."""

    def __init__(self, ws, required, start_row=1, plate_key=None):
        self.ws, self.cols, self.header_row = ws, {}, None
        req = [norm(r) for r in required]
        for r in range(start_row, min(ws.max_row, start_row + 400) + 1):
            row = {norm(ws.cell(r, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None}
            if all(any(k == q or k.startswith(q) for k in row) for q in req):
                self.header_row = r
                for k, c in row.items():
                    self.cols[k] = c
                break
        self.plate_key = norm(plate_key) if plate_key else None

    @property
    def found(self):
        return self.header_row is not None

    def col(self, name):
        n = norm(name)
        if n in self.cols:
            return self.cols[n]
        for k, c in self.cols.items():
            if k.startswith(n):
                return c
        return None

    def rows(self, stop_at_section=True, max_blank=3):
        """Sinh (số dòng, dict giá trị theo tên cột chuẩn hóa)."""
        if not self.found:
            return
        blank = 0
        pc = self.col(self.plate_key) if self.plate_key else None
        for r in range(self.header_row + 1, self.ws.max_row + 1):
            first = txt(self.ws.cell(r, 2).value)
            if stop_at_section and first and first[0] in "①②③④⑤⑥":
                return
            vals = {k: self.ws.cell(r, c).value for k, c in self.cols.items()}
            key_val = self.ws.cell(r, pc).value if pc else None
            if pc and (key_val is None or txt(key_val) == ""):
                blank += 1
                if blank >= max_blank:
                    return
                continue
            blank = 0
            yield r, vals

    def get(self, vals, name):
        n = norm(name)
        if n in vals:
            return vals[n]
        for k, v in vals.items():
            if k.startswith(n):
                return v
        return None


# ---------------------------------------------------------------- dịch vụ


class FleetSheetService:
    def __init__(self, sheet_id=SHEET_ID):
        self.sheet_id = sheet_id
        self._cache = None          # (thời điểm, dữ liệu đã phân tích)
        self._lock = threading.Lock()
        self.last_error = None

    # -- tải file
    def _download(self) -> bytes:
        url = f"https://docs.google.com/spreadsheets/d/{self.sheet_id}/export?format=xlsx"
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=25) as resp:
            return resp.read()

    def load(self, force=False, today=None):
        """Trả dữ liệu đã phân tích (cache CACHE_SECONDS). Lỗi tải mà có cache cũ -> trả cache cũ + cờ."""
        with self._lock:
            now = time.time()
            if not force and self._cache and now - self._cache[0] < CACHE_SECONDS:
                return self._cache[1]
            try:
                data = self.parse(self._download(), today=today)
                self._cache = (now, data)
                self.last_error = None
                return data
            except Exception as e:
                self.last_error = f"{type(e).__name__}: {e}"[:300]
                if self._cache:
                    stale = dict(self._cache[1])
                    stale["stale_error"] = self.last_error
                    return stale
                raise

    def invalidate(self):
        self._cache = None

    # -- phân tích
    def parse(self, xlsx_bytes: bytes, today=None) -> dict:
        today = today or date.today()
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
        tabs = {}
        for ws in wb.worksheets:
            n = norm(ws.title)
            if n.startswith("LUU TRU"):
                continue
            for key, prefix in (("dm", "DM DAU"), ("dodau", "DO DAU"), ("giayto", "GIAY TO"),
                                ("baoduong", "BAO DUONG"), ("vo", "THAY VO"), ("bat", "THAY BAT"),
                                ("tongquan", "TONG QUAN")):
                if n.startswith(prefix) and key not in tabs:
                    tabs[key] = ws
        missing = [k for k in ("dm", "dodau", "giayto", "baoduong", "vo", "bat") if k not in tabs]

        dm = self._parse_dm(tabs.get("dm"))
        fuel = self._parse_fuel(tabs.get("dodau"), today)
        docs, thresholds = self._parse_docs(tabs.get("giayto"), today)
        maint = self._parse_maint(tabs.get("baoduong"), today, fuel["odo_by_plate"])
        tires = self._parse_log(tabs.get("vo"), "vo", today)
        tarps = self._parse_log(tabs.get("bat"), "bat", today)
        debts = self._debts(fuel["rows"], maint["log"], tires, tarps)
        costs = self._monthly_costs(fuel["rows"], maint["log"], tires, tarps, fuel["odo_points"])
        todo = self._todo(docs, maint, fuel, today)

        return {
            "source": {"sheet_id": self.sheet_id, "read_at": datetime.now().strftime("%H:%M:%S %d/%m/%Y"),
                       "tabs": {k: v.title for k, v in tabs.items()}, "missing_tabs": missing},
            "thresholds": thresholds,
            "dm": dm,
            "docs": docs,
            "fuel": {k: v for k, v in fuel.items() if k not in ("odo_points",)},
            "maint": maint,
            "tires": tires,
            "tarps": tarps,
            "debts": debts,
            "costs": costs,
            "todo": todo,
        }

    # ---- DM DẦU (danh mục ô chọn)
    def _parse_dm(self, ws):
        out = {"plates": [], "stations": [], "tank": [], "payment": [], "garages": [], "maint_items": []}
        if ws is None:
            return out
        t = Table(ws, ["SO XE", "CAY DAU", "TINH TRANG BINH"])
        if not t.found:
            return out
        mapping = {"plates": "SO XE", "stations": "CAY DAU", "tank": "TINH TRANG BINH",
                   "payment": "HINH THUC THANH TOAN", "garages": "GARA / NCC", "maint_items": "HANG MUC BAO DUONG"}
        for key, h in mapping.items():
            c = t.col(h)
            if not c:
                continue
            seen = []
            for r in range(t.header_row + 1, ws.max_row + 1):
                v = txt(ws.cell(r, c).value)
                if not v:
                    if r > t.header_row + 3 and not txt(ws.cell(r + 1, c).value):
                        break
                    continue
                if v not in seen:
                    seen.append(v)
            out[key] = seen
        return out

    # ---- GIẤY TỜ XE + ĐĂNG KIỂM
    @staticmethod
    def _thresholds(ws):
        yellow, red = 30, 14
        if ws is None:
            return {"yellow": yellow, "red": red}
        for r in range(1, 12):
            for c in range(1, ws.max_column + 1):
                n = norm(ws.cell(r, c).value)
                if n.startswith("VANG") and to_num(ws.cell(r, c + 1).value):
                    yellow = int(to_num(ws.cell(r, c + 1).value))
                if n.startswith("DO <") or n.startswith("DO ≤") or n == "DO ≤":
                    if to_num(ws.cell(r, c + 1).value):
                        red = int(to_num(ws.cell(r, c + 1).value))
        return {"yellow": yellow, "red": red}

    @staticmethod
    def _status_days(days, th):
        if days is None:
            return "missing"
        if days < 0:
            return "overdue"
        if days <= th["red"]:
            return "soon"
        if days <= th["yellow"]:
            return "near"
        return "ok"

    def _parse_docs(self, ws, today):
        th = self._thresholds(ws)
        docs = []
        if ws is None:
            return docs, th

        def add(kind, plate, trailer, d, note, row):
            dd = to_date(d)
            days = (dd - today).days if dd else None
            docs.append({"kind": kind, "plate": fmt_plate(plate), "trailer": trailer or None,
                         "subject": trailer if "romoóc" in kind and trailer else fmt_plate(plate),
                         "expiry": iso(dd), "days_left": days, "status": self._status_days(days, th),
                         "note": txt(note), "sheet_row": row})

        t = Table(ws, ["SO XE", "SO ROMOOC", "HAN GDD DAU KEO", "HAN GDD ROMOOC"], plate_key="SO XE")
        for r, v in t.rows():
            add("Giấy đi đường – đầu kéo", txt(t.get(v, "SO XE")), txt(t.get(v, "SO ROMOOC")),
                t.get(v, "HAN GDD DAU KEO"), t.get(v, "GHI CHU"), r)
            add("Giấy đi đường – romoóc", txt(t.get(v, "SO XE")), txt(t.get(v, "SO ROMOOC")),
                t.get(v, "HAN GDD ROMOOC"), t.get(v, "GHI CHU"), r)
        t = Table(ws, ["SO XE", "HAN TEM DUONG BO"], plate_key="SO XE")
        for r, v in t.rows():
            add("Tem đường bộ", txt(t.get(v, "SO XE")), None, t.get(v, "HAN TEM DUONG BO"), t.get(v, "GHI CHU"), r)
        t = Table(ws, ["SO XE DAU KEO", "HAN DK DAU KEO", "SO ROMOOC", "HAN DK ROMOOC"], plate_key="SO XE DAU KEO")
        for r, v in t.rows():
            add("Đăng kiểm – đầu kéo", txt(t.get(v, "SO XE DAU KEO")), txt(t.get(v, "SO ROMOOC")),
                t.get(v, "HAN DK DAU KEO"), t.get(v, "GHI CHU"), r)
            add("Đăng kiểm – romoóc", txt(t.get(v, "SO XE DAU KEO")), txt(t.get(v, "SO ROMOOC")),
                t.get(v, "HAN DK ROMOOC"), t.get(v, "GHI CHU"), r)
        return docs, th

    # ---- ĐỔ DẦU
    def _parse_fuel(self, ws, today):
        res = {"rows": [], "consumption": [], "by_plate": {}, "last_entry_date": None,
               "odo_by_plate": {}, "odo_points": defaultdict(list), "issues": []}
        if ws is None:
            return res
        t = Table(ws, ["NGAY DO", "SO XE", "CAY DAU", "SO LIT"], plate_key="SO XE")
        uuid_col = t.col("MA APP")
        rows = []
        for r, v in t.rows(stop_at_section=False, max_blank=5):
            d = to_date(t.get(v, "NGAY DO"))
            liters = to_num(t.get(v, "SO LIT"))
            price = to_num(t.get(v, "DON GIA"))
            amount = to_num(t.get(v, "THANH TIEN")) or (round(liters * price) if liters and price else None)
            note = txt(t.get(v, "GHI CHU"))
            rows.append({
                "sheet_row": r, "uuid": txt(ws.cell(r, uuid_col).value) if uuid_col else "",
                "date": iso(d), "plate": fmt_plate(txt(t.get(v, "SO XE"))),
                "station": txt(t.get(v, "CAY DAU")), "liters": liters, "price": price, "amount": amount,
                "odo": to_num(t.get(v, "KM LUC DO")), "tank": txt(t.get(v, "TINH TRANG BINH")),
                "invoice": txt(t.get(v, "SO HOA DON")), "payment": txt(t.get(v, "THANH TOAN")),
                "paid_date": iso(to_date(t.get(v, "NGAY THANH TOAN"))), "note": note,
                "cancelled": "DA HUY" in norm(note), "flags": [],
            })
        live = [x for x in rows if not x["cancelled"]]
        # --- kiểm tra từng dòng
        last_odo = {}
        for x in sorted(live, key=lambda x: (x["plate"], x["date"] or "", x["odo"] or 0)):
            d = date.fromisoformat(x["date"]) if x["date"] else None
            if d is None:
                x["flags"].append("Thiếu ngày đổ")
            elif d > today:
                x["flags"].append("Ngày ở tương lai")
            elif d.year < 2020 or d < today - timedelta(days=730):
                x["flags"].append("Ngày nghi sai năm")
            if x["liters"] is None or x["liters"] <= 0:
                x["flags"].append("Thiếu số lít")
            elif x["liters"] > 700:
                x["flags"].append("Số lít bất thường (> 700 L)")
            if x["odo"] is None:
                if "TACH" not in norm(x["note"]):
                    x["flags"].append("Thiếu ODO")
            else:
                prev = last_odo.get(x["plate"])
                if prev is not None and x["odo"] < prev:
                    x["flags"].append(f"ODO lùi (trước {prev:,.0f})".replace(",", "."))
                last_odo[x["plate"]] = max(prev or 0, x["odo"])
            if "ODO NGHI SAI" in norm(x["note"]):
                x["flags"].append("ODO ghi chú nghi sai")
            if not x["payment"]:
                x["flags"].append("Chưa ghi thanh toán")
        # --- gộp theo (xe, ngày): HĐ tách 2 lần cùng lượt đổ tính là 1 lần
        groups = defaultdict(lambda: {"liters": 0.0, "odo": None, "full": False, "bad_odo": False})
        for x in live:
            if not x["date"]:
                continue
            g = groups[(x["plate"], x["date"])]
            g["liters"] += x["liters"] or 0
            if x["odo"] is not None and not any(f.startswith("ODO") for f in x["flags"]):
                g["odo"] = max(g["odo"] or 0, x["odo"])
            g["full"] = g["full"] or norm(x["tank"]) == "DAY BINH"
        per_plate = defaultdict(list)
        for (p, d), g in groups.items():
            per_plate[p].append((d, g))
            if g["odo"]:
                res["odo_points"][p].append((d, g["odo"]))
        # --- L/100km giữa 2 lần đầy bình
        for p, items in per_plate.items():
            items.sort()
            prev_full, acc = None, 0.0
            for d, g in items:
                acc += g["liters"]
                if g["full"] and g["odo"]:
                    if prev_full:
                        km = g["odo"] - prev_full[1]
                        liters_window = acc - prev_full[2]
                        if 300 <= km <= 6000:  # dưới 300 km sai số đong đầy bình quá lớn
                            l100 = round(liters_window / km * 100, 1)
                            dev = (l100 - FUEL_NORM) / FUEL_NORM
                            res["consumption"].append({
                                "plate": p, "from": prev_full[0], "to": d, "km": km,
                                "liters": round(liters_window, 1), "l_per_100km": l100,
                                "deviation_pct": round(dev * 100, 1), "alert": abs(dev) > FUEL_DEVIATION})
                    prev_full = (d, g["odo"], acc)
        for p, pts in res["odo_points"].items():
            d, odo = max(pts, key=lambda x: x[1])
            res["odo_by_plate"][p] = {"odo": odo, "date": d}
        res["rows"] = rows
        res["last_entry_date"] = max((x["date"] for x in live if x["date"] and x["date"] <= today.isoformat()), default=None)
        for p in {x["plate"] for x in live}:
            cons = [c for c in res["consumption"] if c["plate"] == p]
            res["by_plate"][p] = {
                "fills": sum(1 for x in live if x["plate"] == p),
                "latest_l_per_100km": cons[-1]["l_per_100km"] if cons else None,
                "avg_l_per_100km": round(sum(c["liters"] for c in cons) / sum(c["km"] for c in cons) * 100, 1) if cons else None,
                "odo": res["odo_by_plate"].get(p),
            }
        res["issues"] = [{"sheet_row": x["sheet_row"], "plate": x["plate"], "date": x["date"], "flags": x["flags"]}
                         for x in live if x["flags"]]
        return res

    # ---- BẢO DƯỠNG
    def _parse_maint(self, ws, today, odo_by_plate):
        out = {"oil": [], "grease": [], "gearbox": [], "log": []}
        if ws is None:
            return out
        th_days = {"red": 14, "yellow": 30}

        def current_km(plate, sheet_km):
            o = (odo_by_plate.get(plate) or {}).get("odo")
            vals = [x for x in (sheet_km, o) if x]
            return max(vals) if vals else None

        t = Table(ws, ["SO XE", "DINH MUC", "KM THAY GAN NHAT"], plate_key="SO XE")
        for r, v in t.rows():
            p = fmt_plate(txt(t.get(v, "SO XE")))
            normkm = to_num(t.get(v, "DINH MUC")) or 15000
            last = to_num(t.get(v, "KM THAY GAN NHAT"))
            cur = current_km(p, to_num(t.get(v, "KM HIEN TAI")))
            remain = (last + normkm - cur) if (last and cur) else None
            warning = None
            if last and cur and cur < last:
                warning = "KM hiện tại nhỏ hơn KM lần thay gần nhất — sổ dầu chưa cập nhật"
                remain = None
            st = "missing" if remain is None else ("overdue" if remain < 0 else "soon" if remain <= OIL_SOON_KM
                                                    else "near" if remain <= OIL_NEAR_KM else "ok")
            out["oil"].append({"plate": p, "norm_km": normkm, "last_km": last,
                               "last_date": iso(to_date(t.get(v, "NGAY BAO DUONG"))), "current_km": cur,
                               "odo_source_date": (odo_by_plate.get(p) or {}).get("date"),
                               "due_km": (last + normkm) if last else None, "remaining_km": remain,
                               "status": st, "warning": warning, "note": txt(t.get(v, "GHI CHU")), "sheet_row": r})
        t = Table(ws, ["SO XE", "NGAY BOI MO GAN NHAT"], plate_key="SO XE")
        for r, v in t.rows():
            last = to_date(t.get(v, "NGAY BOI MO GAN NHAT"))
            due = date(last.year + 1, last.month, min(last.day, 28)) if last else None
            days = (due - today).days if due else None
            out["grease"].append({"plate": fmt_plate(txt(t.get(v, "SO XE"))), "last_date": iso(last),
                                  "due_date": iso(due), "days_left": days,
                                  "status": self._status_days(days, th_days), "sheet_row": r})
        t = Table(ws, ["SO XE", "NGAY THAY GAN NHAT", "KM LUC THAY"], plate_key="SO XE")
        for r, v in t.rows():
            p = fmt_plate(txt(t.get(v, "SO XE")))
            last_d = to_date(t.get(v, "NGAY THAY GAN NHAT"))
            last_km = to_num(t.get(v, "KM LUC THAY"))
            normkm = to_num(t.get(v, "DINH MUC")) or 50000
            cur = current_km(p, to_num(t.get(v, "KM HIEN TAI")))
            remain_km = (last_km + normkm - cur) if (last_km and cur) else None
            due_d = date(last_d.year + 1, last_d.month, min(last_d.day, 28)) if last_d else None
            days = (due_d - today).days if due_d else None
            if remain_km is None and days is None:
                st = "missing"
            else:
                st_km = "ok" if remain_km is None else ("overdue" if remain_km < 0 else "soon" if remain_km <= OIL_SOON_KM
                                                        else "near" if remain_km <= OIL_NEAR_KM else "ok")
                st_d = self._status_days(days, th_days) if days is not None else "ok"
                order = ["overdue", "soon", "near", "ok"]
                st = min(st_km, st_d, key=order.index)
            out["gearbox"].append({"plate": p, "last_date": iso(last_d), "last_km": last_km, "norm_km": normkm,
                                   "current_km": cur, "remaining_km": remain_km, "due_date": iso(due_d),
                                   "days_left": days, "status": st, "sheet_row": r})
        out["log"] = self._parse_log(ws, "baoduong", today)
        return out

    # ---- NHẬT KÝ (bảo dưỡng / vỏ / bạt)
    LOG_SPECS = {
        "baoduong": (["NGAY", "BIEN SO", "HANG MUC", "KM ODO"], "SO TIEN", "GARA / NCC"),
        "vo": (["NGAY", "BIEN SO", "VI TRI", "SO LUONG"], "SO TIEN", "GARA / NCC"),
        "bat": (["NGAY", "BIEN SO", "NOI DUNG", "NCC"], "SO TIEN", "NCC"),
    }

    def _parse_log(self, ws, kind, today):
        if ws is None:
            return []
        required, money_col, vendor_col = self.LOG_SPECS[kind]
        t = Table(ws, required, plate_key="BIEN SO")
        uuid_col = t.col("MA APP")
        out = []
        for r, v in t.rows(stop_at_section=False, max_blank=5):
            note = txt(t.get(v, "GHI CHU"))
            d = to_date(t.get(v, "NGAY"))
            item = {"sheet_row": r, "uuid": txt(ws.cell(r, uuid_col).value) if uuid_col else "",
                    "date": iso(d), "plate": fmt_plate(txt(t.get(v, "BIEN SO"))),
                    "vendor": txt(t.get(v, vendor_col)), "invoice": txt(t.get(v, "SO HD")),
                    "amount": to_num(t.get(v, money_col)), "payment": txt(t.get(v, "THANH TOAN")),
                    "note": note, "cancelled": "DA HUY" in norm(note)}
            if kind == "baoduong":
                item.update({"item": txt(t.get(v, "HANG MUC")), "odo": to_num(t.get(v, "KM ODO")),
                             "content": txt(t.get(v, "NOI DUNG"))})
            elif kind == "vo":
                item.update({"position": txt(t.get(v, "VI TRI")), "quantity": to_num(t.get(v, "SO LUONG")),
                             "brand": txt(t.get(v, "LOAI / HIEU"))})
            else:
                item.update({"content": txt(t.get(v, "NOI DUNG"))})
            out.append(item)
        return out

    # ---- CÔNG NỢ NHÀ CUNG CẤP
    @staticmethod
    def _debts(fuel_rows, maint_log, tires, tarps):
        groups = defaultdict(lambda: {"unpaid": 0.0, "unpaid_count": 0, "paid": 0.0, "no_status": 0.0, "no_status_count": 0, "items": []})

        def add(cat, vendor, x, unpaid, paid):
            g = groups[(cat, vendor or "(chưa ghi NCC)")]
            amt = x.get("amount") or 0
            if unpaid:
                g["unpaid"] += amt
                g["unpaid_count"] += 1
                g["items"].append({"kind": cat, "sheet_row": x["sheet_row"], "uuid": x.get("uuid"), "date": x["date"],
                                   "plate": x["plate"], "amount": amt, "invoice": x.get("invoice")})
            elif paid:
                g["paid"] += amt
            else:
                g["no_status"] += amt
                g["no_status_count"] += 1

        for x in fuel_rows:
            if x["cancelled"]:
                continue
            p = norm(x["payment"])
            add("Dầu", x["station"], x, p.startswith("CONG NO"), p.startswith("DA TT"))
        for cat, rows in (("Bảo dưỡng – gara", maint_log), ("Vỏ xe", tires), ("Bạt", tarps)):
            for x in rows:
                if x["cancelled"]:
                    continue
                p = norm(x["payment"])
                add(cat, x["vendor"], x, p.startswith("CHUA THANH TOAN"), p.startswith("DA THANH TOAN"))
        out = [{"category": c, "vendor": v, **{k: (round(val) if isinstance(val, float) else val) for k, val in g.items()}}
               for (c, v), g in groups.items()]
        out.sort(key=lambda x: (-x["unpaid"], x["category"]))
        return out

    # ---- CHI PHÍ THEO XE THEO THÁNG
    @staticmethod
    def _monthly_costs(fuel_rows, maint_log, tires, tarps, odo_points):
        table = defaultdict(lambda: {"fuel": 0.0, "fuel_liters": 0.0, "maint": 0.0, "tires": 0.0, "tarps": 0.0})
        for x in fuel_rows:
            if not x["cancelled"] and x["date"]:
                c = table[(x["plate"], x["date"][:7])]
                c["fuel"] += x["amount"] or 0
                c["fuel_liters"] += x["liters"] or 0
        for key, rows in (("maint", maint_log), ("tires", tires), ("tarps", tarps)):
            for x in rows:
                if not x["cancelled"] and x["date"]:
                    table[(x["plate"], x["date"][:7])][key] += x["amount"] or 0
        out = []
        for (p, m), c in sorted(table.items()):
            pts = sorted(odo_points.get(p, []))
            in_month = [o for d, o in pts if d[:7] == m]
            before = [o for d, o in pts if d[:7] < m]
            km = None
            if in_month:
                start = max(before) if before else min(in_month)
                km = max(in_month) - start if max(in_month) > start else None
            total = c["fuel"] + c["maint"] + c["tires"] + c["tarps"]
            out.append({"plate": p, "month": m, **{k: round(v) for k, v in c.items()}, "total": round(total),
                        "km": km, "cost_per_km": round(total / km) if km else None})
        return out

    # ---- VIỆC CẦN XỬ LÝ HÔM NAY
    @staticmethod
    def _todo(docs, maint, fuel, today):
        sev = {"overdue": 0, "soon": 1, "near": 2}
        label = {"overdue": "Quá hạn", "soon": "Sắp hạn", "near": "Gần hạn"}
        items = []
        for d in docs:
            if d["status"] in sev:
                items.append({"level": d["status"], "level_label": label[d["status"]], "category": d["kind"],
                              "subject": d["subject"], "plate": d["plate"], "due": d["expiry"],
                              "detail": f"còn {d['days_left']} ngày" if d["days_left"] >= 0 else f"quá {-d['days_left']} ngày"})
        for o in maint["oil"]:
            if o["status"] in sev:
                items.append({"level": o["status"], "level_label": label[o["status"]], "category": "Thay nhớt máy",
                              "subject": o["plate"], "plate": o["plate"], "due": None,
                              "detail": f"còn {o['remaining_km']:,.0f} km".replace(",", ".") if o["remaining_km"] >= 0
                              else f"quá {-o['remaining_km']:,.0f} km".replace(",", ".")})
        for g in maint["grease"] + maint["gearbox"]:
            if g["status"] in sev:
                cat = "Bôi mỡ romoóc" if "days_left" in g and "last_km" not in g else "Thay nhớt hộp số"
                items.append({"level": g["status"], "level_label": label[g["status"]], "category": cat,
                              "subject": g["plate"], "plate": g["plate"], "due": g.get("due_date"),
                              "detail": f"còn {g['days_left']} ngày" if g.get("days_left") is not None else ""})
        items.sort(key=lambda x: (sev[x["level"]], x["due"] or "9999"))
        missing = []

        def miss(cat, plates):
            if plates:
                missing.append({"category": cat, "count": len(plates), "plates": plates})

        miss("Giấy tờ chưa có ngày hạn", sorted({f"{d['subject']} ({d['kind']})" for d in docs if d["status"] == "missing"}))
        miss("Thay nhớt máy thiếu số liệu", [o["plate"] for o in maint["oil"] if o["status"] == "missing"])
        miss("Bôi mỡ romoóc chưa có ngày", [g["plate"] for g in maint["grease"] if g["status"] == "missing"])
        miss("Nhớt hộp số chưa có số liệu", [g["plate"] for g in maint["gearbox"] if g["status"] == "missing"])
        no_odo = [i for i in fuel["issues"] if "Thiếu ODO" in i["flags"]]
        miss("Sổ dầu: dòng thiếu ODO", [f"dòng {i['sheet_row']} ({i['plate']})" for i in no_odo])
        bad = [i for i in fuel["issues"] if any(f.startswith(("ODO lùi", "Ngày", "Số lít")) for f in i["flags"])]
        miss("Sổ dầu: dòng nghi sai (ODO lùi / ngày / số lít)", [f"dòng {i['sheet_row']} ({i['plate']}): {', '.join(i['flags'])}" for i in bad])
        last = fuel.get("last_entry_date")
        lag = (today - date.fromisoformat(last)).days if last else None
        return {
            "counts": {k: sum(1 for i in items if i["level"] == k) for k in sev},
            "missing_count": sum(m["count"] for m in missing),
            "items": items,
            "missing": missing,
            "fuel_log_last_date": last,
            "fuel_log_lag_days": lag,
            "fuel_alerts": [c for c in fuel["consumption"] if c["alert"]][-30:],
        }


fleet_sheet = FleetSheetService()
