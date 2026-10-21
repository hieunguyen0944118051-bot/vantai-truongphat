"""Mục "Xe & Chi phí" (GĐ2): đọc từ Google Sheets, ghi qua Apps Script.

Đọc : GET  /api/xe/data            toàn bộ dữ liệu đã phân tích (cache 3 phút, ?refresh=1 để đọc lại ngay)
      GET  /api/xe/vehicle/{plate} hồ sơ 1 xe
      GET  /api/xe/export?month=YYYY-MM  file Excel cho kế toán
Ghi : POST /api/xe/{kind}          kind = do-dau | bao-duong | thay-vo | thay-bat
      PUT  /api/xe/{kind}/{uuid}   sửa dòng do app ghi
      POST /api/xe/{kind}/{uuid}/huy   hủy (đánh dấu ĐÃ HỦY, không xóa)
      POST /api/xe/cong-no/da-tt   đánh dấu đã thanh toán
"""
import io
import os
import uuid as uuidlib
from datetime import date, datetime
from datetime import date as Date  # dùng cho chú thích kiểu khi trường cũng tên 'date'
from typing import List, Optional

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field

import auth
import models
import sheet_writer
from fleet import FLEET_BY_CLEAN, clean_plate
from fleet_sheet import fleet_sheet, norm

router = APIRouter(prefix="/api/xe", tags=["Xe & Chi phí"])

KIND_TAB = {"do-dau": "dodau", "bao-duong": "baoduong", "thay-vo": "vo", "thay-bat": "bat"}


def write_roles():
    return [r.strip() for r in os.environ.get("XE_WRITE_ROLES", "admin,manager").split(",") if r.strip()]


def can_write(user: models.User = Depends(auth.get_current_user)) -> models.User:
    if user.role not in write_roles():
        raise HTTPException(status_code=403, detail="Tài khoản này chỉ được xem, không được nhập/sửa.")
    return user


def _data(refresh=False):
    try:
        return fleet_sheet.load(force=refresh)
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Không đọc được file Google Sheets xe: {e}")


# ------------------------------------------------------------------ ĐỌC

@router.get("/data")
def get_data(refresh: bool = False, current_user: models.User = Depends(auth.get_current_user)):
    d = _data(refresh)
    return {**d, "write_enabled": sheet_writer.configured(), "can_write": current_user.role in write_roles()}


@router.post("/refresh")
def refresh(current_user: models.User = Depends(auth.get_current_user)):
    fleet_sheet.invalidate()
    d = _data(True)
    return {"ok": True, "read_at": d["source"]["read_at"]}


@router.get("/vehicle/{plate}")
def vehicle_profile(plate: str, current_user: models.User = Depends(auth.get_current_user)):
    d = _data()
    cp = clean_plate(plate)
    info = FLEET_BY_CLEAN.get(cp)
    if not info:
        raise HTTPException(status_code=404, detail="Không có xe này trong danh sách")
    p = info["plate_number"]
    trailer_clean = clean_plate(info.get("trailer_number") or "")

    def mine(x):
        return clean_plate(x.get("plate", "")) == cp

    docs = [x for x in d["docs"] if mine(x)]
    return {
        "vehicle": info,
        "source": d["source"],
        "docs": docs,
        "oil": next((x for x in d["maint"]["oil"] if mine(x)), None),
        "grease": next((x for x in d["maint"]["grease"] if mine(x)), None),
        "gearbox": next((x for x in d["maint"]["gearbox"] if mine(x)), None),
        "maint_log": [x for x in d["maint"]["log"] if mine(x) and not x["cancelled"]][-30:],
        "tires": [x for x in d["tires"] if mine(x) and not x["cancelled"]][-30:],
        "tarps": [x for x in d["tarps"] if mine(x) and not x["cancelled"]][-30:],
        "fuel": [x for x in d["fuel"]["rows"] if mine(x) and not x["cancelled"]][-40:],
        "fuel_summary": d["fuel"]["by_plate"].get(p),
        "consumption": [x for x in d["fuel"]["consumption"] if mine(x)][-20:],
        "costs": [x for x in d["costs"] if mine(x)],
        "todo": [x for x in d["todo"]["items"] if mine(x)],
        "trailer_clean": trailer_clean,
    }


@router.get("/export")
def export_month(month: str = Query(..., pattern=r"^\d{4}-\d{2}$"),
                 current_user: models.User = Depends(auth.get_current_user)):
    """Excel 1 tháng cho kế toán: đổ dầu, bảo dưỡng, vỏ, bạt, công nợ, chi phí theo xe."""
    d = _data()
    wb = openpyxl.Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F2937")

    def sheet(title, headers, rows, money_cols=()):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append(r)
        for col in money_cols:
            for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
                for c in cell:
                    c.number_format = '#,##0'
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, min(40, len(h) + 4))
        ws.freeze_panes = "A2"

    def dmy(s):
        return datetime.strptime(s, "%Y-%m-%d").strftime("%d/%m/%Y") if s else ""

    inm = lambda x: (x.get("date") or "")[:7] == month and not x.get("cancelled")
    sheet("Đổ dầu", ["Ngày", "Số xe", "Cây dầu", "Số lít", "Đơn giá", "Thành tiền", "ODO", "Tình trạng bình", "Số HĐ", "Thanh toán", "Ngày TT", "Ghi chú", "Dòng sheet"],
          [[dmy(x["date"]), x["plate"], x["station"], x["liters"], x["price"], x["amount"], x["odo"], x["tank"], x["invoice"], x["payment"], dmy(x["paid_date"]), x["note"], x["sheet_row"]]
           for x in d["fuel"]["rows"] if inm(x)], money_cols=(5, 6))
    sheet("Bảo dưỡng", ["Ngày", "Biển số", "Hạng mục", "KM ODO", "Nội dung", "Gara/NCC", "Số HĐ", "Số tiền", "Thanh toán", "Ghi chú", "Dòng sheet"],
          [[dmy(x["date"]), x["plate"], x["item"], x["odo"], x["content"], x["vendor"], x["invoice"], x["amount"], x["payment"], x["note"], x["sheet_row"]]
           for x in d["maint"]["log"] if inm(x)], money_cols=(8,))
    sheet("Thay vỏ", ["Ngày", "Biển số", "Vị trí", "Số lượng", "Loại/Hiệu", "Gara/NCC", "Số HĐ", "Số tiền", "Thanh toán", "Ghi chú", "Dòng sheet"],
          [[dmy(x["date"]), x["plate"], x["position"], x["quantity"], x["brand"], x["vendor"], x["invoice"], x["amount"], x["payment"], x["note"], x["sheet_row"]]
           for x in d["tires"] if inm(x)], money_cols=(8,))
    sheet("Thay bạt", ["Ngày", "Biển số", "Nội dung", "NCC", "Số HĐ", "Số tiền", "Thanh toán", "Ghi chú", "Dòng sheet"],
          [[dmy(x["date"]), x["plate"], x["content"], x["vendor"], x["invoice"], x["amount"], x["payment"], x["note"], x["sheet_row"]]
           for x in d["tarps"] if inm(x)], money_cols=(6,))
    sheet("Chi phí theo xe", ["Số xe", "Tháng", "Dầu (đ)", "Số lít", "Bảo dưỡng (đ)", "Vỏ (đ)", "Bạt (đ)", "Tổng (đ)", "Km (theo ODO sổ dầu)", "Chi phí/km (đ)"],
          [[x["plate"], x["month"], x["fuel"], x["fuel_liters"], x["maint"], x["tires"], x["tarps"], x["total"], x["km"], x["cost_per_km"]]
           for x in d["costs"] if x["month"] == month], money_cols=(3, 5, 6, 7, 8, 10))
    sheet("Công nợ (đến hiện tại)", ["Loại", "Nhà cung cấp", "Còn nợ (đ)", "Số dòng nợ", "Đã trả (đ)", "Chưa ghi TT (đ)"],
          [[x["category"], x["vendor"], x["unpaid"], x["unpaid_count"], x["paid"], x["no_status"]] for x in d["debts"]], money_cols=(3, 5, 6))
    del wb["Sheet"]
    info = wb.create_sheet("Nguồn", 0)
    info.append(["Báo cáo tháng", month])
    info.append(["Nguồn", f"Google Sheets {d['source']['sheet_id']}"])
    info.append(["Đọc lúc", d["source"]["read_at"]])
    info.append(["Người xuất", current_user.full_name or current_user.username])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"XE_CHI_PHI_{month}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


# ------------------------------------------------------------------ KIỂM TRA DỮ LIỆU NHẬP

class FuelIn(BaseModel):
    uuid: Optional[str] = None
    date: date
    plate: str
    station: str
    liters: float = Field(gt=0)
    price: float = Field(gt=0)
    odo: Optional[float] = None
    tank: str
    invoice: Optional[str] = ""
    payment: Optional[str] = ""
    paid_date: Optional[date] = None
    note: Optional[str] = ""
    confirm_warnings: bool = False


class MaintIn(BaseModel):
    uuid: Optional[str] = None
    date: date
    plate: str
    item: str
    odo: Optional[float] = None
    content: Optional[str] = ""
    vendor: Optional[str] = ""
    invoice: Optional[str] = ""
    amount: Optional[float] = None
    payment: Optional[str] = ""
    note: Optional[str] = ""
    confirm_warnings: bool = False


class TireIn(BaseModel):
    uuid: Optional[str] = None
    date: date
    plate: str
    position: Optional[str] = ""
    quantity: float = Field(gt=0)
    brand: Optional[str] = ""
    vendor: Optional[str] = ""
    invoice: Optional[str] = ""
    amount: Optional[float] = None
    payment: Optional[str] = ""
    note: Optional[str] = ""
    confirm_warnings: bool = False


class TarpIn(BaseModel):
    uuid: Optional[str] = None
    date: date
    plate: str
    content: str
    vendor: Optional[str] = "Phùng Lĩnh"
    invoice: Optional[str] = ""
    amount: Optional[float] = None
    payment: Optional[str] = ""
    note: Optional[str] = ""
    confirm_warnings: bool = False


def _check_common(body, d, errors, warnings):
    today = date.today()
    if body.date > today:
        errors.append("Ngày ở tương lai")
    if body.date.year < 2020:
        errors.append("Ngày trước năm 2020 — nghi gõ sai năm")
    elif (today - body.date).days > 120:
        warnings.append(f"Ngày cách hôm nay {(today - body.date).days} ngày")
    plates = d["dm"]["plates"]
    if plates and clean_plate(body.plate) not in {clean_plate(p) for p in plates}:
        errors.append("Số xe không có trong danh mục (DM DẦU)")


def _validate_fuel(body: FuelIn, d):
    errors, warnings = [], []
    _check_common(body, d, errors, warnings)
    dm = d["dm"]
    if dm["stations"] and body.station not in dm["stations"]:
        errors.append("Cây dầu không có trong danh mục (DM DẦU)")
    if dm["tank"] and body.tank not in dm["tank"]:
        errors.append("Tình trạng bình không hợp lệ")
    if body.payment and dm["payment"] and body.payment not in dm["payment"]:
        errors.append("Hình thức thanh toán không có trong danh mục")
    if body.liters > 700:
        errors.append("Số lít > 700 L — vượt dung tích bình")
    elif body.liters < 20:
        warnings.append("Số lít rất ít (< 20 L)")
    if not 10000 <= body.price <= 40000:
        warnings.append("Đơn giá ngoài khoảng 10.000–40.000 đ/lít")
    if body.odo is None:
        if body.tank == "Đầy bình":
            warnings.append("Đầy bình nhưng chưa nhập ODO — không tính được L/100km")
    else:
        prev = (d["fuel"]["by_plate"].get(_fmt(body.plate)) or {}).get("odo")
        if prev and body.odo < prev["odo"]:
            errors.append(f"ODO lùi: lần trước {prev['odo']:,.0f} km ngày {prev['date']}".replace(",", "."))
        elif prev and body.odo - prev["odo"] > 8000:
            warnings.append(f"ODO tăng {body.odo - prev['odo']:,.0f} km so với lần trước — kiểm tra lại".replace(",", "."))
    return errors, warnings


def _fmt(p):
    return FLEET_BY_CLEAN.get(clean_plate(p), {}).get("plate_number", p)


def _validate_simple(body, d):
    errors, warnings = [], []
    _check_common(body, d, errors, warnings)
    if getattr(body, "amount", None) is not None and body.amount < 0:
        errors.append("Số tiền âm")
    if getattr(body, "amount", None) in (None, 0):
        warnings.append("Chưa nhập số tiền")
    if body.payment and body.payment not in ("Đã thanh toán", "Chưa thanh toán"):
        errors.append("Thanh toán chỉ nhận 'Đã thanh toán' / 'Chưa thanh toán'")
    if isinstance(body, MaintIn):
        items = d["dm"]["maint_items"]
        if items and body.item not in items:
            errors.append("Hạng mục không có trong danh mục (DM DẦU)")
    vendors = d["dm"]["garages"]
    if getattr(body, "vendor", "") and vendors and body.vendor not in vendors:
        warnings.append("Gara/NCC chưa có trong danh mục — sheet sẽ báo ô chọn không hợp lệ")
    return errors, warnings


def _plate_raw(p):
    return clean_plate(p)  # sheet ghi biển số dạng 63E01132


def _values(kind, body):
    """Đổi dữ liệu form -> {tiêu đề cột chuẩn hóa: giá trị} đúng tên cột trong sheet."""
    if kind == "dodau":
        return {"NGAY DO": body.date.isoformat(), "SO XE": _plate_raw(body.plate), "CAY DAU": body.station,
                "SO LIT": body.liters, "DON GIA": body.price, "KM LUC DO": body.odo if body.odo is not None else "",
                "TINH TRANG BINH": body.tank, "SO HOA DON": body.invoice or "", "THANH TOAN": body.payment or "",
                "NGAY THANH TOAN": body.paid_date.isoformat() if body.paid_date else "", "GHI CHU": body.note or ""}
    if kind == "baoduong":
        return {"NGAY": body.date.isoformat(), "BIEN SO": _plate_raw(body.plate), "HANG MUC": body.item,
                "KM ODO": body.odo if body.odo is not None else "", "NOI DUNG": body.content or "",
                "GARA / NCC": body.vendor or "", "SO HD": body.invoice or "", "SO TIEN": body.amount if body.amount is not None else "",
                "THANH TOAN": body.payment or "", "GHI CHU": body.note or ""}
    if kind == "vo":
        return {"NGAY": body.date.isoformat(), "BIEN SO": _plate_raw(body.plate), "VI TRI": body.position or "",
                "SO LUONG": body.quantity, "LOAI / HIEU": body.brand or "", "GARA / NCC": body.vendor or "",
                "SO HD": body.invoice or "", "SO TIEN": body.amount if body.amount is not None else "",
                "THANH TOAN": body.payment or "", "GHI CHU": body.note or ""}
    return {"NGAY": body.date.isoformat(), "BIEN SO": _plate_raw(body.plate), "NOI DUNG": body.content,
            "NCC": body.vendor or "", "SO HD": body.invoice or "", "SO TIEN": body.amount if body.amount is not None else "",
            "THANH TOAN": body.payment or "", "GHI CHU": body.note or ""}


def _submit(kind_url, body, user, record_uuid=None):
    tab = KIND_TAB[kind_url]
    d = _data()
    errors, warnings = (_validate_fuel(body, d) if tab == "dodau" else _validate_simple(body, d))
    if errors:
        raise HTTPException(status_code=422, detail={"errors": errors, "warnings": warnings})
    if warnings and not body.confirm_warnings:
        raise HTTPException(status_code=409, detail={"errors": [], "warnings": warnings,
                                                     "message": "Kiểm tra lại rồi bấm Xác nhận để ghi"})
    who = user.full_name or user.username
    try:
        if record_uuid:
            out = sheet_writer.send("update", {"tab": tab, "uuid": record_uuid, "values": _values(tab, body)}, who)
        else:
            out = sheet_writer.send("append", {"tab": tab, "uuid": body.uuid or str(uuidlib.uuid4()),
                                               "values": _values(tab, body)}, who)
    except sheet_writer.SheetWriteError as e:
        raise HTTPException(status_code=503, detail=str(e))
    fleet_sheet.invalidate()
    return {"ok": True, "row": out.get("row"), "duplicate": out.get("duplicate", False),
            "message": out.get("message"), "warnings": warnings}


# ------------------------------------------------------------------ GHI

@router.post("/do-dau")
def add_fuel(body: FuelIn, user: models.User = Depends(can_write)):
    return _submit("do-dau", body, user)


@router.post("/bao-duong")
def add_maint(body: MaintIn, user: models.User = Depends(can_write)):
    return _submit("bao-duong", body, user)


@router.post("/thay-vo")
def add_tire(body: TireIn, user: models.User = Depends(can_write)):
    return _submit("thay-vo", body, user)


@router.post("/thay-bat")
def add_tarp(body: TarpIn, user: models.User = Depends(can_write)):
    return _submit("thay-bat", body, user)


@router.put("/do-dau/{record_uuid}")
def edit_fuel(record_uuid: str, body: FuelIn, user: models.User = Depends(can_write)):
    return _submit("do-dau", body, user, record_uuid)


@router.put("/bao-duong/{record_uuid}")
def edit_maint(record_uuid: str, body: MaintIn, user: models.User = Depends(can_write)):
    return _submit("bao-duong", body, user, record_uuid)


@router.put("/thay-vo/{record_uuid}")
def edit_tire(record_uuid: str, body: TireIn, user: models.User = Depends(can_write)):
    return _submit("thay-vo", body, user, record_uuid)


@router.put("/thay-bat/{record_uuid}")
def edit_tarp(record_uuid: str, body: TarpIn, user: models.User = Depends(can_write)):
    return _submit("thay-bat", body, user, record_uuid)


class CancelIn(BaseModel):
    reason: Optional[str] = ""


@router.post("/{kind}/{record_uuid}/huy")
def cancel(kind: str, record_uuid: str, body: CancelIn, user: models.User = Depends(can_write)):
    if kind not in KIND_TAB:
        raise HTTPException(status_code=404, detail="Không biết loại nhật ký")
    try:
        out = sheet_writer.send("cancel", {"tab": KIND_TAB[kind], "uuid": record_uuid, "reason": body.reason or ""},
                                user.full_name or user.username)
    except sheet_writer.SheetWriteError as e:
        raise HTTPException(status_code=503, detail=str(e))
    fleet_sheet.invalidate()
    return {"ok": True, "row": out.get("row"), "message": out.get("message")}


class PaidItem(BaseModel):
    kind: str            # "Dầu" | "Bảo dưỡng – gara" | "Vỏ xe" | "Bạt"
    sheet_row: int
    uuid: Optional[str] = ""
    plate: str
    date: Optional[str] = None
    amount: Optional[float] = None


class PaidIn(BaseModel):
    items: List[PaidItem]
    method: str = "transfer"   # transfer | cash (chỉ áp dụng dòng dầu)
    date: Optional[Date] = None


CATEGORY_TAB = {"Dầu": "dodau", "Bảo dưỡng – gara": "baoduong", "Vỏ xe": "vo", "Bạt": "bat"}


@router.post("/cong-no/da-tt")
def mark_paid(body: PaidIn, user: models.User = Depends(can_write)):
    if not body.items:
        raise HTTPException(status_code=422, detail="Chưa chọn dòng nào")
    if body.date and body.date > date.today():
        raise HTTPException(status_code=422, detail="Ngày thanh toán ở tương lai")
    items = []
    for it in body.items:
        tab = CATEGORY_TAB.get(it.kind)
        if not tab:
            raise HTTPException(status_code=422, detail=f"Không biết loại công nợ: {it.kind}")
        items.append({"tab": tab, "sheet_row": it.sheet_row, "uuid": it.uuid or "", "plate": clean_plate(it.plate),
                      "date": it.date, "amount": it.amount})
    try:
        out = sheet_writer.send("mark_paid", {"items": items, "method": body.method,
                                              "date": (body.date or date.today()).isoformat()},
                                user.full_name or user.username)
    except sheet_writer.SheetWriteError as e:
        raise HTTPException(status_code=503, detail=str(e))
    fleet_sheet.invalidate()
    return {"ok": True, "rows": out.get("rows", []), "skipped": out.get("skipped", []), "message": out.get("message")}
